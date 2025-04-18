from pypdf import PdfReader # pdf
from ffmpeg import FFmpeg # audio/video
import xml.etree.ElementTree as ET  # xml
from PIL import Image, ExifTags
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
import magic
import datetime
import odfmeta as odf

def extract_metadata_pdf(in_file) -> dict:
	reader = PdfReader(in_file)
	return reader.metadata

# audio and video
def extract_metadata_media(in_file) -> dict:
	ffprobe = FFmpeg(executable="ffprobe")
	ffprobe.input(in_file, print_format="xml", show_streams=None)

	metadata = ffprobe.execute()
	metadata_string = metadata.decode() # Convert Bytes to String

	tree = ET.fromstring(metadata_string) # Parse xml tree
	root = tree[0][0] # Set root node i.e (ffprobe -> Streams -> Stream)
	# 														Root--^
	return root.attrib

def extract_metadata_raster_image(in_file) -> dict:
	image = Image.open(in_file)
	exif = image._getexif()
	if exif == None:
		return {}

	exif = {
		ExifTags.TAGS[k]: v
		for k, v in exif.items()
		if k in ExifTags.TAGS
	}

	# Convert Degree notation to decimal
	if "GPSInfo" in exif.keys():
		gps_tag = exif["GPSInfo"]
		N = gps_tag[2][0] + gps_tag[2][1]/60.0 + gps_tag[2][2]/3600.0
		E = gps_tag[4][0] + gps_tag[4][1]/60.0 + gps_tag[4][2]/3600.0
		exif["GPSInfo"] = f"{format(N, '.4f')} N, {format(E, '.4f')} E" # Overwrite

	metadata = {k: v for k, v in exif.items() if v} if exif else {}
	return metadata

def extract_metadata_svg(in_file) -> dict:
	tree = ET.parse(in_file)
	root = tree.getroot()
	return root.attrib

def extract_metadata_xlsx(in_file) -> dict:
	workbook = load_workbook(in_file)
	properties = workbook.properties
	metadata = {
		'creator': properties.creator,
		'title': properties.title,
		'description': properties.description,
		'subject': properties.subject,
		'identifier': properties.identifier,
		'language': properties.language,
		'created': properties.created,
		'modified': properties.modified,
		'lastModifiedBy': properties.lastModifiedBy,
		'category': properties.category,
		'contentStatus': properties.contentStatus,
		'version': properties.version,
		'revision': properties.revision,
		'keywords': properties.keywords,
		'lastPrinted': properties.lastPrinted
	}
	return metadata

def extract_metadata_docx(in_file) -> dict:
	document = Document(in_file)
	properties = document.core_properties
	metadata = {
		'title': properties.title,
		'subject': properties.subject,
		'identifier': properties.identifier,
		'language': properties.language,
		'created': properties.created,
		'modified': properties.modified,
		'last_modified_by': properties.last_modified_by,
		'category': properties.category,
		'content_status': properties.content_status,
		'version': properties.version,
		'revision': properties.revision,
		'keywords': properties.keywords,
		'last_printed': properties.last_printed,
		'comments': properties.comments
	}
	return metadata

def extract_metadata_pptx(in_file) -> dict:
	presentation = Presentation(in_file)
	properties = presentation.core_properties
	metadata = {
		'title': properties.title,
		'subject': properties.subject,
		'identifier': properties.identifier,
		'language': properties.language,
		'created': properties.created,
		'modified': properties.modified,
		'last_modified_by': properties.last_modified_by,
		'category': properties.category,
		'content_status': properties.content_status,
		'version': properties.version,
		'revision': properties.revision,
		'keywords': properties.keywords,
		'last_printed': properties.last_printed,
		'comments': properties.comments
	}
	return metadata

def print_data(dictionary) -> None:
	if dictionary == None:
		print("No metadata found!")
		return
	items = dictionary.items()
	for k, v in items:
		print(k, ":", v)

def extract_metadata_file(file: str):
	file_type = magic.from_file(file, mime=True)
	metadata = {}

	match file_type:
		case "image/jpeg" | "image/png":
			metadata = extract_metadata_raster_image(file)
		case "application/pdf":
			metadata = extract_metadata_pdf(file)
		case "image/svg+xml":
			metadata = extract_metadata_svg(file)
		case "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
			metadata = extract_metadata_xlsx(file)
		case "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
			metadata = extract_metadata_docx(file)
		case "application/vnd.openxmlformats-officedocument.presentationml.presentation":
			metadata = extract_metadata_pptx(file)
		case _:
			if "video/" in file_type or "audio/" in file_type:
				metadata = extract_metadata_media(file)
			if "application/vnd.oasis.opendocument" in file_type:
				metadata = extract_metadata_odf(file)
	return metadata

def add_meta_metadata(metadata: dict, file: str) -> dict:
	metadata['Meta-Metadata'] = {
		'Metadata Creator': 'Zeta Data Extractor',
		'Metadata Creation Date': datetime.datetime.now().isoformat(),
		'Metadata Last Modified By': 'Zeta Data Extractor',
		'Metadata Last Modified Date': datetime.datetime.now().isoformat(),
		'Metadata Version': '1.0',
		'Metadata Status': 'Complete',
		'Metadata Standard': 'Custom',
		'Metadata Language': 'en',
		'Metadata Source': file,
		'Metadata Accuracy': 'High',
		'Metadata Completeness': 'Full',
		'Metadata Coverage': 'All',
		'Metadata Accessibility': 'Public',
		'Metadata Rights': 'Open',
		'Metadata Relationship': 'None'
	}
	return metadata

def extract_metadata_odf(in_file):
	odf_file = odf.OdfMeta(in_file)
	return odf_file.Properties()
